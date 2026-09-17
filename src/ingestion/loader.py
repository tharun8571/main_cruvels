"""
Document ingestion: accepts a small set of legal documents (PDF/TXT/DOCX),
extracts usable text, and preserves document/page metadata for later
citation. Falls back to PaddleOCR-VL per-page when direct text extraction
comes back empty (i.e. scanned pages).
"""
from __future__ import annotations

import logging
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf as fitz  # use modern pymupdf API

from .ocr import ocr_image
from src.config import get_settings

logger = logging.getLogger(__name__)


@dataclass
class PageContent:
    page_number: int
    text: str
    source_type: str  # "native" or "ocr"
    confidence: float = 1.0


@dataclass
class LoadedDocument:
    doc_id: str
    file_name: str
    pages: list[PageContent] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())


def _load_pdf(path: Path) -> LoadedDocument:
    settings = get_settings()
    min_conf = settings["ingestion"]["min_confidence"]

    doc = fitz.open(path)
    pages: list[PageContent] = []

    for i, page in enumerate(doc):
        native_text = page.get_text("text").strip()

        if native_text:
            pages.append(PageContent(page_number=i + 1, text=native_text, source_type="native"))
            continue

        # No extractable text -> render page to image and OCR it.
        logger.info("Page %s of %s has no native text, running OCR", i + 1, path.name)
        ocr_text = ""
        ocr_conf = 0.0
        temp_img = Path(tempfile.gettempdir()) / f"page_ocr_{i}_{path.stem}.png"
        try:
            pix = page.get_pixmap(dpi=200)
            pix.save(str(temp_img))
            ocr_result = ocr_image(str(temp_img), page_number=i + 1)
            ocr_text = ocr_result.text
            ocr_conf = ocr_result.confidence
        except Exception as ocr_err:
            logger.warning("OCR failed on page %s of %s: %s", i + 1, path.name, ocr_err)
        finally:
            if temp_img.exists():
                try:
                    temp_img.unlink()
                except Exception:
                    pass

        if ocr_conf < min_conf:
            logger.warning(
                "OCR confidence %.2f below threshold %.2f on page %s of %s",
                ocr_conf, min_conf, i + 1, path.name,
            )

        pages.append(
            PageContent(
                page_number=i + 1,
                text=ocr_text,
                source_type="ocr",
                confidence=ocr_conf,
            )
        )

    doc.close()
    return LoadedDocument(doc_id=path.stem, file_name=path.name, pages=pages)


def _load_txt(path: Path) -> LoadedDocument:
    text = path.read_text(encoding="utf-8", errors="ignore")
    page = PageContent(page_number=1, text=text, source_type="native")
    return LoadedDocument(doc_id=path.stem, file_name=path.name, pages=[page])


def _load_docx(path: Path) -> LoadedDocument:
    from docx import Document as DocxDocument

    docx_doc = DocxDocument(path)
    text = "\n".join(p.text for p in docx_doc.paragraphs)
    page = PageContent(page_number=1, text=text, source_type="native")
    return LoadedDocument(doc_id=path.stem, file_name=path.name, pages=[page])


_LOADERS = {
    ".pdf": _load_pdf,
    ".txt": _load_txt,
    ".docx": _load_docx,
}


def load_document(file_path: str | Path) -> LoadedDocument:
    path = Path(file_path)
    settings = get_settings()
    supported = settings["ingestion"]["supported_extensions"]

    if path.suffix.lower() not in supported:
        raise ValueError(f"Unsupported file type: {path.suffix}. Supported: {supported}")

    loader_fn = _LOADERS[path.suffix.lower()]
    logger.info("Loading document: %s", path.name)
    return loader_fn(path)
